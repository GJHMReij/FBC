"""
Prints the numbers needed for the selection flowchart (no modelling, runs in
seconds). Runs the same selection and split as pe_pipeline.py for Model 1/2 and
for Model 3 (scans without D-dimer excluded before the split).

Run in Spyder on MyDRE with F5, then copy-paste the printed block.
"""
import argparse
import io
import json
import sys
from contextlib import redirect_stdout
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pe_pipeline as pp

INPUT_CSV = r"C:\Users\Max.Reijers\Desktop\CohortMLgeslacht.csv"
OUTCOME_CORRECTION = r"C:\Users\Max.Reijers\Desktop\df_met_script8000.xlsx"
FAKE_DATES = False   # True only for a test on the dummy data on the Mac


def numbers(models, require_ddimer):
    args = argparse.Namespace(
        input_csv=Path(INPUT_CSV), outcome_correction=Path(OUTCOME_CORRECTION),
        models=models, require_ddimer=require_ddimer, fake_dates_for_testing=FAKE_DATES,
    )
    with redirect_stdout(io.StringIO()):   # keep the output short
        df, y, _, _ = pp.load_and_prepare(args)
        is_test, cutoff, n_both, n_late = pp.temporal_split(df["scan_date"], pp.patient_ids(df), 0.2)
    train, test = ~is_test, is_test
    out = dict(df.attrs["selection"])
    out.update({
        "n_train": int(train.sum()), "pe_train": int(y[train].sum()),
        "n_test": int(test.sum()), "pe_test": int(y[test].sum()),
        "test_share_pct": round(100 * test.mean(), 2),
        "patients_in_both_periods": int(n_both),
        "pe_total": int(y.sum()), "pe_pct": round(100 * y.mean(), 1),
        "first_scan_date": str(df["scan_date"].min().date()),
        "last_scan_date": str(df["scan_date"].max().date()),
        "n_patients": int(pp.patient_ids(df).nunique()),
    })
    if pp.HOSPITAL_COL in df.columns:
        out["hospital_train"] = df.loc[train, pp.HOSPITAL_COL].value_counts().to_dict()
        out["hospital_test"] = df.loc[test, pp.HOSPITAL_COL].value_counts().to_dict()
    return out


res12 = numbers([1], False)
res3 = numbers([3], True)

for title, res in [("MODELS 1 AND 2 (all scans)", res12),
                   ("MODEL 3 (scans without D-dimer excluded before the split)", res3)]:
    print("=" * 60)
    print(title)
    print("=" * 60)
    for k, v in res.items():
        print(f"{k}: {v}")
    print()

# ---- document with the same numbers (opens in Word; HTML if python-docx is missing)
out_dir = Path(pp.REPO_ROOT) / "results_temporal"
out_dir.mkdir(parents=True, exist_ok=True)
with open(out_dir / "flowchart_numbers.json", "w") as f:
    json.dump({"models_1_2": res12, "model_3": res3}, f, indent=1, default=str)

sections = [("Models 1 and 2 (all scans)", res12),
            ("Model 3 (scans without D-dimer excluded before the split)", res3)]
try:
    import docx
    doc = docx.Document()
    doc.add_heading("Numbers for the selection flowchart", level=0)
    for title, res in sections:
        doc.add_heading(title, level=1)
        t = doc.add_table(rows=1, cols=2)
        t.style = "Light Grid Accent 1"
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Item", "Value"
        for k, v in res.items():
            cells = t.add_row().cells
            cells[0].text, cells[1].text = str(k), str(v)
    path = out_dir / "flowchart_numbers.docx"
    doc.save(str(path))
except ImportError:
    html = ["<html><head><meta charset='utf-8'></head><body><h1>Numbers for the selection flowchart</h1>"]
    for title, res in sections:
        html.append(f"<h2>{title}</h2><table border=1>")
        html += [f"<tr><td>{k}</td><td>{v}</td></tr>" for k, v in res.items()]
        html.append("</table>")
    html.append("</body></html>")
    path = out_dir / "flowchart_numbers.html"
    path.write_text("\n".join(html), encoding="utf-8")
print(f"Document written: {path}")
print("\nCopy everything above and paste it in the chat.")
