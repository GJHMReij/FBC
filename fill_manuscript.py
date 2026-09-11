"""
Fill the manuscript's Table 2 (performance metrics) and Table 3 (feature
importance ranking) with Model 1/2/3's internal-validation results, using
the *_manuscript_data.json files each model script saves as a side effect
of its normal run (no manual copy-pasting from the console needed). Also
appends the combined ROC curve (from combine_roc_curves.py) at the end of
the document, with a label -- inserting it precisely next to its existing
Figure 2 caption is skipped deliberately (the template's paragraphs mix
multiple figure captions together in ways that make automatic placement
there fragile/risky to get right blind); move it into place manually in
Word, that's a quick manual step.

Works fine with only 1 or 2 of the three models run so far -- it just
leaves the other rows/columns blank, same as the original template.

Run with the project venv, ideally after combine_roc_curves.py:
    .venv/bin/python fill_manuscript.py "path/to/CBC PE manuscript - empty.docx"
"""
import argparse
import json
from pathlib import Path

import docx
from docx.shared import Inches

SCRIPT_DIR = Path(__file__).resolve().parent

MODEL_DATA_FILES = [
    SCRIPT_DIR / "model1_manuscript_data.json",
    SCRIPT_DIR / "model2_manuscript_data.json",
    SCRIPT_DIR / "model3_manuscript_data.json",
]

# Table 2 (doc.tables[1]): row index of each model's "Internal validation" row.
TABLE2_INTERNAL_ROWS = [2, 5, 8]  # Model 1, 2, 3
# Table 3 (doc.tables[2]): rows 2-11 are ranks 1-10; columns 1/2/3 are the
# internal-validation columns for Model 1/2/3 (columns 4-6 are external/UCLH,
# left untouched -- no external data yet).
TABLE3_RANK_ROW_START = 2
TABLE3_MODEL_COLS = [1, 2, 3]


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("manuscript_path", type=Path, help="Path to the manuscript .docx to fill")
    parser.add_argument(
        "--output", type=Path, default=None,
        help="Output path (default: '<manuscript name> - filled.docx' next to the input)",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    output_path = args.output or args.manuscript_path.with_name(
        args.manuscript_path.stem + " - filled.docx"
    )

    model_data = []
    for path in MODEL_DATA_FILES:
        if not path.exists():
            print(f"Skipping {path.name} (not found -- run that model first if you want it included)")
            model_data.append(None)
            continue
        with open(path) as f:
            model_data.append(json.load(f))

    doc = docx.Document(str(args.manuscript_path))
    table2 = doc.tables[1]
    table3 = doc.tables[2]

    for data, row_idx in zip(model_data, TABLE2_INTERNAL_ROWS):
        if data is None:
            continue
        row = table2.rows[row_idx]
        row.cells[1].text = f"{data['auc']:.3f}"
        row.cells[2].text = f"{data['sensitivity_95']:.3f}"
        row.cells[3].text = f"{data['specificity_95']:.3f}"
        row.cells[4].text = f"{data['accuracy_95']:.3f}"
        row.cells[5].text = f"{data['f1_95']:.3f}"
        row.cells[6].text = f"{data['brier']:.3f}"
        print(f"Filled Table 2 row {row_idx} ({data['label']})")

    for data, col_idx in zip(model_data, TABLE3_MODEL_COLS):
        if data is None:
            continue
        for rank, (name, importance) in enumerate(data["top_features"]):
            row = table3.rows[TABLE3_RANK_ROW_START + rank]
            row.cells[col_idx].text = name
        print(f"Filled Table 3 column {col_idx} ({data['label']})")

    roc_path = SCRIPT_DIR / "combined_roc_curve.png"
    if roc_path.exists():
        doc.add_page_break()
        doc.add_heading("Figure 2: ROC curves (internal validation) -- move into place manually", level=2)
        doc.add_picture(str(roc_path), width=Inches(5.5))
        print(f"Appended {roc_path.name} at the end (move it next to the Figure 2 caption manually)")
    else:
        print(f"Skipping {roc_path.name} (not found -- run combine_roc_curves.py first)")

    doc.save(str(output_path))
    print(f"\nFilled manuscript saved to: {output_path}")


if __name__ == "__main__":
    main()
