"""
Word report with all results of one model (written next to the other outputs).

Uses python-docx. If python-docx is not installed, an HTML file with the same
content is written instead (Word opens .html files too) and a warning is printed.
"""
from pathlib import Path

import pandas as pd


def _ci(d):
    return f"{d['est']:.3f} ({d['lo']:.3f}-{d['hi']:.3f})"


def _table_rows(tbl):
    cols = [tbl.index.name or ""] + [str(c) for c in tbl.columns]
    rows = [[str(i)] + [("" if pd.isna(v) else (f"{v:.3f}" if isinstance(v, float) else str(v)))
                        for v in r] for i, r in zip(tbl.index, tbl.values.tolist())]
    return cols, rows


def build_report(path_prefix, info):
    """info: dict with label, settings (list of (name, value)), split (list of str),
    test_metrics (list of (name, formatted)), tables (list of (title, DataFrame, note)),
    hospitals (list of rows), importance (Series), images (list of (title, path))."""
    path_prefix = str(path_prefix)
    try:
        import docx
        from docx.shared import Inches
    except ImportError:
        return _build_html(path_prefix, info)

    doc = docx.Document()
    doc.add_heading(f"{info['label']} - results overview", level=0)
    doc.add_paragraph(info.get("intro", ""))

    def add_table(cols, rows):
        t = doc.add_table(rows=1, cols=len(cols))
        t.style = "Light Grid Accent 1"
        for j, c in enumerate(cols):
            t.rows[0].cells[j].text = str(c)
        for r in rows:
            cells = t.add_row().cells
            for j, v in enumerate(r):
                cells[j].text = str(v)

    doc.add_heading("1. Data and split", level=1)
    for line in info["split"]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("2. Model settings", level=1)
    add_table(["Setting", "Value"], [[a, b] for a, b in info["settings"]])

    doc.add_heading("3. Main performance (hold-out test set, Rubin-pooled)", level=1)
    add_table(["Metric", "Estimate (95% CI)"], [[a, b] for a, b in info["test_metrics"]])

    doc.add_heading("4. Figures", level=1)
    for title, p in info["images"]:
        if Path(p).exists():
            doc.add_paragraph(title)
            doc.add_picture(str(p), width=Inches(4.8))

    n = 5
    for title, tbl, note in info["tables"]:
        doc.add_heading(f"{n}. {title}", level=1)
        if note:
            doc.add_paragraph(note)
        cols, rows = _table_rows(tbl)
        add_table(cols, rows)
        n += 1

    if info.get("hospitals"):
        doc.add_heading(f"{n}. Test set per hospital", level=1)
        add_table(info["hospitals"][0], info["hospitals"][1:])
        n += 1

    doc.add_heading(f"{n}. Feature importance (top 20)", level=1)
    imp = info["importance"].head(20)
    add_table(["Feature", "Importance"], [[i, f"{v:.4f}"] for i, v in imp.items()])

    out = f"{path_prefix}_report.docx"
    doc.save(out)
    return out


def _build_html(path_prefix, info):
    print("WARNING: python-docx not installed -> writing an .html report instead "
          "(install with: pip install python-docx)")
    h = [f"<html><head><meta charset='utf-8'><title>{info['label']}</title></head><body>",
         f"<h1>{info['label']} - results overview</h1><p>{info.get('intro', '')}</p>",
         "<h2>1. Data and split</h2><ul>" + "".join(f"<li>{x}</li>" for x in info["split"]) + "</ul>",
         "<h2>2. Model settings</h2><table border=1>" +
         "".join(f"<tr><td>{a}</td><td>{b}</td></tr>" for a, b in info["settings"]) + "</table>",
         "<h2>3. Main performance (hold-out test set, Rubin-pooled)</h2><table border=1>" +
         "".join(f"<tr><td>{a}</td><td>{b}</td></tr>" for a, b in info["test_metrics"]) + "</table>",
         "<h2>4. Figures</h2>"]
    for title, p in info["images"]:
        h.append(f"<p>{title}<br><img src='{Path(p).name}' width='480'></p>")
    for k, (title, tbl, note) in enumerate(info["tables"], start=5):
        cols, rows = _table_rows(tbl)
        h.append(f"<h2>{k}. {title}</h2><p>{note or ''}</p><table border=1><tr>" +
                 "".join(f"<th>{c}</th>" for c in cols) + "</tr>" +
                 "".join("<tr>" + "".join(f"<td>{v}</td>" for v in r) + "</tr>" for r in rows) +
                 "</table>")
    if info.get("hospitals"):
        hs = info["hospitals"]
        h.append("<h2>Test set per hospital</h2><table border=1><tr>" +
                 "".join(f"<th>{c}</th>" for c in hs[0]) + "</tr>" +
                 "".join("<tr>" + "".join(f"<td>{v}</td>" for v in r) + "</tr>" for r in hs[1:]) +
                 "</table>")
    h.append("<h2>Feature importance (top 20)</h2><table border=1>" +
             "".join(f"<tr><td>{i}</td><td>{v:.4f}</td></tr>" for i, v in info["importance"].head(20).items()) +
             "</table></body></html>")
    out = f"{path_prefix}_report.html"
    Path(out).write_text("\n".join(h), encoding="utf-8")
    return out
